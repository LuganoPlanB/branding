"""Generate editable Plan B layouts and a starter deck. Run from repository root."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import argparse
import re
import subprocess
import sys
import shutil
import zipfile
from lxml import etree
import cairosvg
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'slides'
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--theme', choices=['light', 'dark'])
args = parser.parse_args()
if args.theme is None:
    for mode in ('light', 'dark'):
        subprocess.run([sys.executable, __file__, '--theme', mode], check=True)
    # Keep existing download URLs working as aliases of the light edition.
    for stem, ext in [('starter', 'pptx'), ('template', 'potx')]:
        shutil.copyfile(OUT / f'planb-{stem}-light.{ext}', OUT / f'planb-{stem}.{ext}')
    sys.exit(0)
MODE = args.theme
DARK = MODE == 'dark'
T = json.loads((ROOT / 'tokens/tokens.json').read_text())[MODE]
C = {k: T['--planb-color-' + v].lstrip('#') for k,v in {
    'canvas':'canvas','cool':'canvas-alt','ink':'text','navy':'accent-strong',
    'sky':'accent','gold':'highlight','cyan':'neon-cyan',
    'violet':'neon-violet','pink':'neon-pink'}.items()}
C.update(panel='211D1D' if DARK else '082952', onpanel='FFFFFF' if DARK else 'FFFEFA',
         mist='494040' if DARK else 'DCEAF6', header='211D1D' if DARK else '082952')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
for shape in list(prs.slide_master.shapes):
    shape._element.getparent().remove(shape._element)
prs.core_properties.title = f'Plan B — Civic Digital Commons / {MODE.title()}'
prs.core_properties.subject = 'Reusable 16:9 slide layouts and editable starter deck'
prs.core_properties.author = 'Plan B Foundation'
prs.core_properties.keywords = 'Plan B, branding, template, Civic Digital Commons'
prs.core_properties.comments = 'Generated from slides/build_template.py. Examples are instructions or explicitly synthetic.'

def rgb(color): return RGBColor.from_string(C.get(color,color))

def rect(s,x,y,w,h,color,radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color); sh.line.fill.background()
    # Remove Office's default style reference: renderers otherwise add a shadow
    # to the brand bands even when the local effect list is empty.
    style=sh._element.find(qn('p:style'))
    if style is not None: sh._element.remove(style)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    if radius: sh.adjustments[0]=.12
    return sh

def text(s,txt,x,y,w,h,size=22,color='ink',bold=False,ph=None,name=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    sh.name=name or (f'Placeholder {ph[1]}' if ph else txt[:40])
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Inter'; p.font.size=Pt(size); p.font.bold=bold
        p.alignment=PP_ALIGN.LEFT
        ppr=p._p.get_or_add_pPr(); ppr.set('marL','0'); ppr.set('indent','0'); ppr.append(OxmlElement('a:buNone'))
        p.font.color.rgb=rgb(color); p.space_after=Pt(12); p.line_spacing=1.12
    if ph:
        nv=sh._element.find(qn('p:nvSpPr')).find(qn('p:nvPr'))
        el=OxmlElement('p:ph'); el.set('type',ph[0]); el.set('idx',str(ph[1])); nv.append(el)
    return sh

def gradient(s,x,y,w,h):
    sh=rect(s,x,y,w,h,'gold'); sh.name='Plan B / Network energy gradient'; sp=sh._element.spPr
    sp.remove(sp.find(qn('a:solidFill')))
    g=OxmlElement('a:gradFill'); g.set('rotWithShape','1'); stops=OxmlElement('a:gsLst')
    for col, percent in re.findall(r'(#[0-9a-fA-F]{6}) (\d+)%', T['--planb-gradient-rainbow']):
        pos=int(percent)*1000
        stop=OxmlElement('a:gs'); stop.set('pos',str(pos)); c=OxmlElement('a:srgbClr'); c.set('val',col.lstrip('#').upper()); stop.append(c); stops.append(stop)
    g.append(stops); lin=OxmlElement('a:lin'); lin.set('ang','0'); lin.set('scaled','1'); g.append(lin); sp.insert_element_before(g,'a:ln','a:effectLst')
    return sh

def logo(s,x,y,w,variant='white'):
    path=ROOT/f'assets/logos/lugano_planb_2026_logo_wide_{variant}.svg'
    raw=path.read_bytes(); fallback=cairosvg.svg2png(bytestring=raw,output_width=1416)
    pic=s.shapes.add_picture(BytesIO(fallback),Inches(x),Inches(y),width=Inches(w))
    pic.name="Original Lugano's Plan B mark"
    pic._element.xpath('.//p:cNvPr')[0].set('descr',"Lugano's Plan B original logo; existing mark, not a new Foundation logo")
    part=Part(PackURI(f'/ppt/media/planb-{variant}.svg'),'image/svg+xml',prs.part.package,raw)
    rid=s.part.relate_to(part,RT.IMAGE)
    blip=pic._element.xpath('.//a:blip')[0]; exts=OxmlElement('a:extLst'); ext=OxmlElement('a:ext'); ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    svg=etree.Element('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip'); svg.set(qn('r:embed'),rid)
    ext.append(svg); exts.append(ext); blip.append(ext)
    return pic

def chrome(s,dark=False):
    bg=s.background.fill; bg.solid(); bg.fore_color.rgb=rgb('cool' if dark else 'canvas')
    fg='navy'
    gradient(s,.65,6.74,12,.10)
    text(s,'Plan B Foundation',.65,6.98,4,.25,11,fg,True)
    text(s,'Civic Digital Commons',9.4,6.98,3.25,.25,11,fg)

def heading(s,title,sub=None):
    text(s,title,.65,.55,12,1.1,36,'navy',True,('title',0))
    if sub: text(s,sub,.65,1.65,11.7,.75,19,'ink',False,('body',1))

layouts=list(prs.slide_layouts)
for layout in layouts:
    for sh in list(layout.shapes): sh._element.getparent().remove(sh._element)
    layout._element.set('showMasterSp','1')
    layout._element.set('type','cust')
    layout._element.set('preserve','1')

def new(i,name,dark=False):
    layout=layouts[i]; layout.name=f'Plan B / {MODE.title()} / '+name
    s=prs.slides.add_slide(layout); s._element.set('showMasterSp','1'); chrome(s,dark)
    return s

def finalize(s,notes):
    """Place reusable geometry in the layout, keeping editable content on the slide."""
    layout=s.slide_layout
    layout._element.cSld.insert(0,deepcopy(s._element.cSld.bg))
    for sh in list(s.shapes):
        cloned=deepcopy(sh._element)
        for node in cloned.iter():
            for attr,val in list(node.attrib.items()):
                if attr in [qn('r:embed'),qn('r:link')]:
                    rel=s.part.rels[val]
                    node.set(attr,layout.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,is_external=rel.is_external))
        layout.shapes._spTree.insert_element_before(cloned,'p:extLst')
        if not sh.is_placeholder: sh._element.getparent().remove(sh._element)
    s.notes_slide.notes_text_frame.text=notes

s=new(0,'Cover')
text(s,'[Presentation\ntitle]',.65,1.35,7.3,2.15,48,'navy',True,('title',0))
text(s,'[A clear sentence about the civic purpose]',.68,3.9,6.8,.9,24,'ink',False,('subTitle',1))
text(s,'[Presenter / team]\n[Event · date]',.68,5.55,6.5,.8,17,'ink',False,('body',2))
rect(s,8.55,0,4.79,6.65,'panel'); logo(s,9.08,1.45,3.72)
text(s,'Open.\nCapable.\nCurious.\nGrounded.',9.08,3.05,3.6,2.3,29,'onpanel',True)
gradient(s,8.55,6.25,4.79,.4)
finalize(s,'COVER. Replace bracketed text. Aim for a title under 8 words. Keep the inherited mark intact. This template uses the existing Lugano Plan B identity, not a newly designed Foundation logo. Install bundled Inter fonts before editing. Choose Home > New Slide for Plan B layouts.')

s=new(1,'Section',True)
text(s,'[Section title]',.8,2.05,11.6,1.7,48,'navy',True,('title',0))
text(s,'[One sentence that sets the direction]',.85,4.15,10.8,.9,25,'ink',False,('body',1))
gradient(s,.85,1.35,3.6,.22)
finalize(s,'SECTION. Use between major topics, not before every slide. The background follows the selected light or dark edition. The gradient is the brand signal. Keep the statement short.')

s=new(2,'Title and content')
heading(s,'[Lead with the point]','[Explain why this matters to the audience]')
text(s,'[First idea — one concise sentence]\n[Second idea — one concise sentence]\n[Third idea — one concise sentence]',.85,2.85,11.3,2.9,27,'ink',False,('body',2))
# The common gradient band carries the brand without competing with body text.
finalize(s,'CONTENT. Use three short statements instead of a wall of bullets. Body is 27pt. Explain public benefit before technical detail. The editable text box supports additional paragraphs but do not shrink below 20pt to fit more content.')

s=new(3,'Two columns')
heading(s,'[Compare two ideas]')
text(s,'[First perspective]',.7,2.05,5.55,.6,26,'navy',True,('body',1))
text(s,'[Explain the idea in 40 words or fewer. Keep sentences concrete and easy to read.]',.7,2.95,5.45,2.8,24,'ink',False,('body',2))
text(s,'[Second perspective]',7.05,2.05,5.55,.6,26,'navy',True,('body',3))
text(s,'[Explain the difference or consequence. Use the same level of detail as the first column.]',7.05,2.95,5.45,2.8,24,'ink',False,('body',4))
rect(s,6.63,2,.012,4.15,'mist')
finalize(s,'TWO COLUMNS. Useful for problem/response, today/proposal, or two options. Equal space means equal status. Do not imply a ranking through color alone.')

s=new(4,'Image and story')
heading(s,'[Make the place visible]')
rect(s,.65,2,7.35,4.25,'cool',True)
text(s,'Insert a documentary photograph',1,3.1,6.6,1.15,26,'navy',True)
text(s,'[Insert a documentary photograph]',.65,2,7.35,4.25,26,'navy',True,('pic',1))
text(s,'[What should we notice?]',8.5,2.2,4.1,1.1,27,'navy',True,('body',2))
text(s,'[Connect the image to the public purpose. Use real Lugano, civic spaces, workshops or people making and learning.]',8.5,3.55,4.1,2.6,22,'ink',False,('body',3))
text(s,'[Photo credit / source / permission]',.7,6.4,11.5,.35,12,'ink',False,('body',4))
finalize(s,'IMAGE. Use the picture placeholder to insert an image, then Crop > Fit or Fill. Replace the placeholder copy. Add alt text and a credit. Use real photographs with permission; no synthetic evidence or implied endorsements.')

s=new(5,'Chart and takeaway')
heading(s,'[Say what the data shows]','[Unit, population and time period]')
text(s,'[Insert chart]',.7,2.5,7.8,3.75,22,'navy',False,('chart',2))
text(s,'[The takeaway]',9,2.65,3.6,.85,27,'navy',True,('body',3))
text(s,'[Explain the implication and the limits of the evidence.]',9,3.75,3.6,2.1,22,'ink',False,('body',4))
text(s,'ILLUSTRATIVE DATA ONLY · Replace values and add a source before presenting.',.7,6.45,11.9,.35,12,'ink',False,('body',5))
finalize(s,'CHART. The starter contains an editable native PowerPoint chart with an embedded Excel workbook. All values are synthetic. Use Edit Data to replace them. Give units, source, date and uncertainty. Preserve the selected light or dark edition colors and keep all labels readable.')
for sh in list(s.shapes):
    if sh.is_placeholder and sh.placeholder_format.idx==2: sh._element.getparent().remove(sh._element)
data=CategoryChartData(); data.categories=['Example A','Example B','Example C']; data.add_series('Illustrative values',[20,35,50])
chart=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(.75),Inches(2.65),Inches(7.8),Inches(3.55),data).chart
chart.has_legend=False; chart.has_title=False
chart.chart_style=10
chart_style=OxmlElement('c:spPr')
chart_fill=OxmlElement('a:solidFill'); chart_color=OxmlElement('a:srgbClr'); chart_color.set('val',C['canvas'])
chart_fill.append(chart_color); chart_style.append(chart_fill)
chart_line=OxmlElement('a:ln'); chart_line.append(OxmlElement('a:noFill')); chart_style.append(chart_line)
chart._chartSpace.insert_element_before(chart_style,'c:txPr','c:externalData','c:printSettings','c:userShapes','c:extLst')
chart.category_axis.tick_labels.font.name='Inter'; chart.category_axis.tick_labels.font.size=Pt(15)
chart.value_axis.tick_labels.font.name='Inter'; chart.value_axis.tick_labels.font.size=Pt(12)
chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=60
chart.plots[0].has_data_labels=True; chart.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
chart.plots[0].data_labels.font.name='Inter'; chart.plots[0].data_labels.font.size=Pt(15)
for labels in (chart.category_axis.tick_labels, chart.value_axis.tick_labels, chart.plots[0].data_labels):
    labels.font.color.rgb=rgb('ink')
chart.value_axis.major_gridlines.format.line.color.rgb=rgb('mist')
for point,col in zip(chart.series[0].points,['violet','sky','gold'] if DARK else ['navy','sky','gold']): point.format.fill.solid(); point.format.fill.fore_color.rgb=rgb(col); point.format.line.fill.background()

s=new(6,'Table')
heading(s,'[Make a decision easier]','[Explain what the comparison includes]')
text(s,'[Insert table]',.7,2.55,11.95,3.6,22,'navy',False,('tbl',2))
text(s,'[Source / definitions / caveat]',.7,6.35,11.9,.4,13,'ink',False,('body',3))
finalize(s,'TABLE. Native editable table in the starter. Keep to 3–4 columns and 3–5 rows. Replace bracketed fields. Do not use a screenshot of a spreadsheet. Keep text at least 20pt. Preserve the selected edition header and alternating quiet grounds.')
for sh in list(s.shapes):
    if sh.is_placeholder and sh.placeholder_format.idx==2: sh._element.getparent().remove(sh._element)
table=s.shapes.add_table(4,3,Inches(.7),Inches(2.55),Inches(11.95),Inches(3.4)).table
rows=[['[Option]','[Public benefit]','[Consideration]'],['[Option A]','[Benefit]','[Constraint]'],['[Option B]','[Benefit]','[Constraint]'],['[Option C]','[Benefit]','[Constraint]']]
for ri,row in enumerate(rows):
    for ci,txt in enumerate(row):
        cell=table.cell(ri,ci); cell.text=txt; cell.margin_left=Inches(.18); cell.margin_top=Inches(.14)
        cell.fill.solid(); cell.fill.fore_color.rgb=rgb('header' if ri==0 else ('cool' if ri%2 else 'canvas'))
        for p in cell.text_frame.paragraphs: p.font.name='Inter'; p.font.size=Pt(20); p.font.bold=ri==0; p.font.color.rgb=rgb('onpanel' if ri==0 else 'ink')

s=new(7,'Quotation')
text(s,'[A short quotation,\naccurately attributed.]',.85,1.6,11.3,2.9,42,'navy',True,('title',0))
text(s,'[Speaker / role]\n[Source · date]',.9,5,10.7,1.1,20,'ink',False,('body',1))
gradient(s,.85,1.05,2.8,.16)
finalize(s,'QUOTE. This is placeholder text, not a real quotation or endorsement. Verify the wording, permission where required, attribution and source. Keep quotations short. Do not invent testimonials.')

s=new(8,'Next steps')
heading(s,'[What happens next?]')
for i,(title,body) in enumerate([('[Learn]','[Concrete next action]'),('[Test]','[Concrete next action]'),('[Contribute]','[Concrete next action]')]):
    x=.7+i*4.15
    text(s,title,x,2.25,3.7,.75,28,'navy',True,('body',1+i*2))
    text(s,body+'\n[Owner · timing]',x,3.4,3.65,2.1,23,'ink',False,('body',2+i*2))
text(s,'[Contact / project URL]',.7,6.1,11.8,.45,19,'navy',False,('body',7))
finalize(s,'NEXT STEPS. Close with specific actions and realistic timing. Add a real contact or project link. The three headings are editable prompts, not a promised programme.')

s=new(9,'Template guide')
heading(s,'Make it yours. Keep it Plan B.','Delete this guidance slide before presenting.')
text(s,'Use the layouts\nInstall Inter, then choose Home → New Slide. Replace bracketed prompts. Duplicate the chart or table example to retain its styling.',.7,2.75,5.6,2.9,23,'ink',False,('body',2))
text(s,'Keep it clear\nOne idea per slide. Keep the selected light or dark palette. Preserve the gradient band and keep text on solid grounds. Aim for 20pt or larger body text.',7,2.75,5.6,2.9,23,'ink',False,('body',3))
for i,col in enumerate(['navy','sky','gold','canvas']):
    rect(s,.7+i*.6,6,.4,.25,col)
text(s,'Full guidance: plan-b.foundation/branding/',3.4,6,8.8,.4,15,'navy')
finalize(s,'GUIDE. This deck is an editable starter, not a finished presentation. Ten Plan B layouts plus a branded blank are supplied. Speaker notes describe usage. Reuse the POTX in PowerPoint to create a new document. Keynote/LibreOffice may vary in rendering; check after import.')

# An additional empty layout for content that does not fit the supplied patterns.
layout=layouts[10]; layout.name=f'Plan B / {MODE.title()} / Blank'; layout.background.fill.solid(); layout.background.fill.fore_color.rgb=rgb('canvas')

band=next(sh for sh in layouts[0].shapes if sh.name=='Plan B / Network energy gradient')
layout.shapes._spTree.insert_element_before(deepcopy(band._element),'p:extLst')

# Apply actual theme colors/fonts so new objects and charts inherit Plan B defaults.
theme=prs.slide_master.part.part_related_by(RT.THEME)
tree=etree.fromstring(theme.blob)
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
scheme=tree.find('.//a:clrScheme',ns); scheme.set('name','Plan B Civic Digital Commons')
mapping={'dk1':'ink','lt1':'canvas','dk2':'navy','lt2':'cool','accent1':'navy','accent2':'sky','accent3':'gold','accent4':'violet','accent5':'pink','accent6':'cyan','hlink':'navy','folHlink':'violet'}
for slot,col in mapping.items():
    el=scheme.find('a:'+slot,ns)
    for child in list(el): el.remove(child)
    value=etree.SubElement(el,qn('a:srgbClr')); value.set('val',C[col])
for el in tree.findall('.//a:fontScheme',ns):
    el.set('name','Plan B Inter')
    for latin in el.findall('.//a:latin',ns): latin.set('typeface','Inter')
theme._blob=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
starter=OUT/f'planb-starter-{MODE}.pptx'
template=OUT/f'planb-template-{MODE}.potx'
prs.save(starter)

# POTX differs from PPTX at the main-part content type, not just the extension.
with zipfile.ZipFile(starter) as src, zipfile.ZipFile(template,'w',zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist():
        blob=src.read(item.filename)
        if item.filename=='[Content_Types].xml': blob=blob.replace(b'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',b'application/vnd.openxmlformats-officedocument.presentationml.template.main+xml')
        dst.writestr(item,blob)
print(f'Generated {MODE}: 10 editable starter slides, 11 named layouts with gradients, PPTX and POTX.')
